"""Word document export for completed code reviews."""
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt, RGBColor


def word_report(report: dict) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.85)
    document.styles['Normal'].font.name = 'Calibri'
    document.styles['Normal'].font.size = Pt(11)
    for name in ('Title', 'Heading 1', 'Heading 2'):
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    document.core_properties.title = 'CodeScanner code review report'
    document.core_properties.author = 'CodeScanner'
    document.add_heading('CodeScanner code review report', 0)
    document.add_paragraph(f"Source: {report['source']}")
    document.add_heading('Summary', 1)
    document.add_paragraph(report['summary'])
    document.add_paragraph(report['notice'])
    document.add_paragraph(
        f"Reviewed {report['scope']['files_reviewed']} files; "
        f"excluded {len(report['scope']['skipped'])}."
    )
    if report['incomplete']:
        document.add_paragraph('Partial review: at least one specialist failed.')
    if report.get('synthesis_fallback'):
        document.add_paragraph('AI synthesis was unavailable. Specialist findings were combined and sorted automatically.')
    document.add_heading('Findings', 1)
    if not report['findings']:
        document.add_paragraph('No findings were reported.')
    for finding in report['findings']:
        document.add_heading(f"[{finding['severity'].upper()}] {finding['title']}", 2)
        document.add_paragraph(f"{finding['category']} · {finding['file']}:{finding['line']}")
        document.add_paragraph(finding['explanation'])
        document.add_paragraph(f"Suggested fix: {finding['recommendation']}")
    document.add_heading('Specialist summaries', 1)
    for agent in report['agents']:
        document.add_heading(agent['agent'], 2)
        document.add_paragraph(agent['summary'])
    if report['scope']['skipped']:
        document.add_heading('Excluded files', 1)
        for skipped in report['scope']['skipped']:
            document.add_paragraph(f"{skipped['path']}: {skipped['reason']}", style='List Bullet')
    output = BytesIO()
    document.save(output)
    return output.getvalue()
